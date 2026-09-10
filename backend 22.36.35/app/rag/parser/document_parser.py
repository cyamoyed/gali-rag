import io


class DocumentParser:
    @staticmethod
    async def parse(file_content: bytes, file_type: str, filename: str) -> str:
        if file_type == "application/pdf":
            return DocumentParser._parse_pdf(file_content)
        elif file_type in (
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ):
            return DocumentParser._parse_docx(file_content)
        elif file_type in (
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ):
            return DocumentParser._parse_excel(file_content)
        elif file_type in ("text/plain", "text/markdown"):
            return file_content.decode("utf-8", errors="ignore")
        elif file_type.startswith("image/"):
            return DocumentParser._parse_image(file_content)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def _parse_pdf(content: bytes) -> str:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts)

    @staticmethod
    def _parse_docx(content: bytes) -> str:
        from docx import Document

        doc = Document(io.BytesIO(content))
        return "\n\n".join(
            para.text for para in doc.paragraphs if para.text.strip()
        )

    @staticmethod
    def _parse_excel(content: bytes) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(
                    str(cell) if cell is not None else "" for cell in row
                )
                if row_text.strip(" |"):
                    text_parts.append(row_text)
        return "\n".join(text_parts)

    @staticmethod
    def _parse_image(content: bytes) -> str:
        try:
            import pytesseract
            from PIL import Image

            image = Image.open(io.BytesIO(content))
            return pytesseract.image_to_string(image, lang="chi_sim+eng")
        except Exception:
            return "[OCR not available]"
